from __future__ import annotations

import hashlib
import os
import shutil
import uuid
import zipfile
from dataclasses import asdict, dataclass
from email import policy
from email.parser import BytesParser
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader

from .models import ActionKind, Artifact, PlannedAction, Task, TaskStatus
from .registry import DesktopToolRegistry
from .service import DesktopTaskService
from .workspace import DesktopWorkspace


class ArtifactBundleError(RuntimeError):
    pass


@dataclass(frozen=True)
class StagedArtifact:
    kind: str
    staged_path: str
    final_path: str
    sha256: str
    summary: str

    def to_preview(self) -> dict[str, object]:
        return {"artifact": asdict(self)}

    @classmethod
    def from_preview(cls, preview: dict[str, object]) -> "StagedArtifact":
        raw = preview.get("artifact")
        if not isinstance(raw, dict):
            raise ArtifactBundleError("交付动作缺少 artifact 预览数据。")
        return cls(
            kind=str(raw["kind"]),
            staged_path=str(raw["staged_path"]),
            final_path=str(raw["final_path"]),
            sha256=str(raw["sha256"]),
            summary=str(raw["summary"]),
        )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class ArtifactBundleDelivery:
    """Deliver an already-verified group of Office artifacts after approval.

    The workflow remains responsible for producing and reviewing the files. This
    class provides the shared last mile: immutable previews, preflight checks,
    confirmation, registry execution, artifacts, and Trace events.
    """

    def __init__(
        self,
        service: DesktopTaskService,
        workspace: DesktopWorkspace,
        registry: DesktopToolRegistry,
    ) -> None:
        self.service = service
        self.workspace = workspace
        self.registry = registry

    def capture(
        self,
        *,
        kind: str,
        staged_path: Path,
        final_filename: str,
        summary: str,
    ) -> StagedArtifact:
        staged_path = staged_path.resolve()
        if not staged_path.is_file():
            raise ArtifactBundleError(f"暂存产物不存在：{staged_path}")
        if not self.workspace.can_write_artifact(staged_path):
            raise ArtifactBundleError(f"暂存产物不在任务工作区内：{staged_path}")
        final_path = (self.workspace.output_root / final_filename).resolve()
        if not self.workspace.is_output_artifact(final_path):
            raise ArtifactBundleError(f"正式交付路径越界：{final_path}")
        return StagedArtifact(
            kind=kind,
            staged_path=str(staged_path),
            final_path=str(final_path),
            sha256=sha256_file(staged_path),
            summary=summary,
        )

    def action(
        self,
        *,
        action_id: str,
        skill: str,
        artifact: StagedArtifact,
    ) -> PlannedAction:
        return PlannedAction(
            action_id=action_id,
            skill=skill,
            kind=ActionKind.WRITE,
            args={"destination": artifact.final_path},
            summary=artifact.summary,
            preview=artifact.to_preview(),
        )

    def confirm_and_deliver(self, task: Task, *, approved: bool) -> Task:
        if task.status is not TaskStatus.AWAITING_CONFIRMATION:
            raise ArtifactBundleError("当前任务不在等待确认状态。")
        self.service.confirm(task, approved=approved)
        if not approved:
            return task

        actions: list[tuple[PlannedAction, StagedArtifact]] = []
        try:
            for action in task.actions:
                artifact = StagedArtifact.from_preview(action.preview)
                self._validate(artifact)
                actions.append((action, artifact))
            if not actions:
                raise ArtifactBundleError("任务没有可交付产物。")

            self.service.trace_store.append(
                task.task_id,
                "artifact_bundle_preflight_passed",
                {
                    "artifact_count": len(actions),
                    "artifacts": [artifact.final_path for _, artifact in actions],
                },
            )

            for action, artifact in actions:
                result = self.registry.execute(
                    task,
                    action,
                    callback=lambda a=artifact: self._commit(a),
                    verify=lambda path, a=artifact: self._verify_committed(path, a),
                )
                self.service.add_artifact(
                    task,
                    Artifact(
                        kind=artifact.kind,
                        staged_path=artifact.staged_path,
                        final_path=str(result),
                        sha256=artifact.sha256,
                        summary=artifact.summary,
                    ),
                )
            self.service.finish(task)
            return task
        except Exception as exc:
            self.service.trace_store.append(
                task.task_id,
                "artifact_bundle_delivery_failed",
                {"error": str(exc)},
            )
            self.service.finish(task, error=str(exc))
            return task

    def _validate(self, artifact: StagedArtifact) -> None:
        staged = Path(artifact.staged_path).resolve()
        final = Path(artifact.final_path).resolve()
        if not staged.is_file() or not self.workspace.can_write_artifact(staged):
            raise ArtifactBundleError(f"暂存产物无效：{staged}")
        if not self.workspace.is_output_artifact(final):
            raise ArtifactBundleError(f"正式交付路径越界：{final}")
        if final.exists():
            raise ArtifactBundleError(f"正式交付文件已存在，拒绝覆盖：{final}")
        if staged.suffix.lower() != final.suffix.lower():
            raise ArtifactBundleError("暂存产物与正式文件扩展名不一致。")
        if sha256_file(staged) != artifact.sha256:
            raise ArtifactBundleError(f"暂存产物哈希已变化：{staged}")
        self._validate_file_structure(staged)

    def _commit(self, artifact: StagedArtifact) -> Path:
        staged = Path(artifact.staged_path).resolve()
        final = Path(artifact.final_path).resolve()
        final.parent.mkdir(parents=True, exist_ok=True)
        temporary = final.with_name(f".{final.name}.{uuid.uuid4().hex}.tmp")
        try:
            shutil.copyfile(staged, temporary)
            if sha256_file(temporary) != artifact.sha256:
                raise ArtifactBundleError(f"复制后的文件哈希不一致：{final}")
            os.rename(temporary, final)
        finally:
            temporary.unlink(missing_ok=True)
        return final

    def _verify_committed(self, path: Path, artifact: StagedArtifact) -> bool:
        path = path.resolve()
        if path != Path(artifact.final_path).resolve() or not path.is_file():
            return False
        if sha256_file(path) != artifact.sha256:
            return False
        try:
            self._validate_file_structure(path)
        except Exception:
            return False
        return True

    @staticmethod
    def _validate_file_structure(path: Path) -> None:
        suffix = path.suffix.lower()
        if path.stat().st_size == 0:
            raise ArtifactBundleError(f"产物为空文件：{path}")
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            if len(reader.pages) < 1:
                raise ArtifactBundleError(f"PDF 没有页面：{path}")
            return
        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                if not workbook.sheetnames:
                    raise ArtifactBundleError(f"XLSX 没有工作表：{path}")
            finally:
                workbook.close()
            return
        if suffix == ".eml":
            message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            if str(message.get("X-LocalDesk-Draft", "")).lower() != "true":
                raise ArtifactBundleError(f"EML 未标记为未发送草稿：{path}")
            return
        if suffix in {".pptx", ".docx"}:
            required = "ppt/presentation.xml" if suffix == ".pptx" else "word/document.xml"
            if not zipfile.is_zipfile(path):
                raise ArtifactBundleError(f"Office 产物不是有效 ZIP 包：{path}")
            with zipfile.ZipFile(path) as archive:
                if required not in archive.namelist():
                    raise ArtifactBundleError(f"Office 产物缺少核心结构：{path}")
            return
        raise ArtifactBundleError(f"暂不支持交付该文件类型：{suffix}")
