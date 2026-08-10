"""Run and confirm the local reimbursement auxiliary demo.

The demo creates small synthetic inputs, uses LibreOffice for the DOCX-to-PDF
render check, stages all results, and never sends email.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook

from localdesk.desktop.docx_delivery import LibreOfficeDocxRenderer
from localdesk.desktop.policy import DesktopPolicyGuard
from localdesk.desktop.reimbursement_workflow import ReimbursementWorkflow
from localdesk.desktop.service import DesktopTaskService
from localdesk.desktop.skills.reimbursement import ReimbursementSkill
from localdesk.desktop.trace_store import TaskTraceStore
from localdesk.desktop.workspace import DesktopWorkspace, WorkspaceConfig


def _workspace(root: Path) -> tuple[DesktopWorkspace, DesktopTaskService]:
    sources, output, tasks = root / "sources", root / "deliveries", root / "tasks"
    for directory in (sources, output, tasks):
        directory.mkdir(parents=True, exist_ok=True)
    workspace = DesktopWorkspace(WorkspaceConfig(
        read_roots=[sources], output_root=output, task_root=tasks,
    ))
    service = DesktopTaskService(DesktopPolicyGuard(workspace), TaskTraceStore(workspace))
    return workspace, service


def _create_inputs(sources: Path) -> tuple[Path, list[Path], Path]:
    payments = sources / "payments.xlsx"
    rules = sources / "rules.docx"
    invoices = [sources / "invoice-001.pdf", sources / "invoice-002.pdf"]
    if not payments.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "payments"
        sheet.append(["payment_id", "invoice_no", "amount"])
        sheet.append(["PAY-001", "INV-001", 50.00])
        sheet.append(["PAY-002", "INV-002", 100.00])
        sheet.append(["PAY-003", "INV-MISSING", 20.00])
        workbook.save(payments)
    if not rules.exists():
        document = Document()
        document.add_heading("Reimbursement rules", level=1)
        document.add_paragraph("max_payment_amount: 80.00")
        document.add_paragraph("amount_tolerance: 0.01")
        document.save(rules)
    for path, invoice_no, amount in (
        (invoices[0], "INV-001", 50.00),
        (invoices[1], "INV-002", 100.00),
    ):
        if path.exists():
            continue
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), f"Invoice No: {invoice_no}\nTotal: {amount:.2f}")
        pdf.save(path)
        pdf.close()
    return payments, invoices, rules


def run_demo(root: Path | None = None, *, main_agent_plan: dict[str, object] | None = None, user_query: str | None = None) -> dict[str, object]:
    root = (root or Path.cwd() / ".localdesk" / "reimbursement-demo-v2").resolve()
    workspace, service = _workspace(root)
    payments, invoices, rules = _create_inputs(root / "sources")
    task = service.create_task(user_query or "核对支付流水、发票和报销规则，生成待复核清单")
    if main_agent_plan is not None:
        service.trace_store.append(task.task_id, "main_agent_planned", main_agent_plan)
    workflow = ReimbursementWorkflow(service, ReimbursementSkill(workspace))
    result = workflow.run(
        task,
        str(payments),
        [str(path) for path in invoices],
        str(rules),
        LibreOfficeDocxRenderer(),
        "finance-review@example.com",
    )
    analysis = result["result"]
    return {
        "task_id": task.task_id,
        "task_status": task.status.value,
        "root": str(root),
        "matched_count": analysis.matched_count,
        "flagged_count": analysis.flagged_count,
        "applied_rules": analysis.applied_rules,
        "staged": {
            key: str(result[key]) for key in ("xlsx", "docx", "pdf", "email", "confirmation")
        },
        "trace": str(workspace.task_dir(task.task_id) / "events.jsonl"),
    }


def confirm_demo(task_id: str, root: Path | None = None, *, approved: bool = True) -> dict[str, object]:
    root = (root or Path.cwd() / ".localdesk" / "reimbursement-demo-v2").resolve()
    workspace, service = _workspace(root)
    task = service.trace_store.load_task_object(task_id)
    workflow = ReimbursementWorkflow(service, ReimbursementSkill(workspace))
    workflow.confirm_and_deliver(task, approved=approved)
    return {
        "task_id": task.task_id,
        "task_status": task.status.value,
        "artifacts": [artifact.final_path for artifact in task.artifacts],
        "trace": str(workspace.task_dir(task.task_id) / "events.jsonl"),
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Run or confirm the LocalDesk reimbursement demo.")
    parser.add_argument("--root", type=Path)
    parser.add_argument("--confirm", metavar="TASK_ID")
    parser.add_argument("--reject", action="store_true")
    args = parser.parse_args()
    payload = confirm_demo(args.confirm, args.root, approved=not args.reject) if args.confirm else run_demo(args.root)
    print(json.dumps(payload, ensure_ascii=False, indent=2))
