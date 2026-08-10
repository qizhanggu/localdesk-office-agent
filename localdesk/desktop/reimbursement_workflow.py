"""将确定性报销核对接入 Task Trace；产物仍停留在 staging，等待统一交付。"""
from __future__ import annotations
from dataclasses import asdict
from pathlib import Path
from docx import Document
from localdesk.desktop.artifact_bundle import ArtifactBundleDelivery
from localdesk.desktop.docx_delivery import DocxRenderer
from localdesk.desktop.skills.office_artifacts import EmailDraftSkill, PdfDeliverySkill, StagedEmailDraft, StagedPdf
from localdesk.desktop.models import ActionKind, PlannedAction, Task, TaskStatus
from localdesk.desktop.registry import DesktopToolRegistry, create_desktop_registry
from localdesk.desktop.service import DesktopTaskService, TaskStateError
from localdesk.desktop.skills.reimbursement import ReimbursementResult, ReimbursementSkill

class ReimbursementWorkflow:
    def __init__(self, service: DesktopTaskService, skill: ReimbursementSkill, registry: DesktopToolRegistry | None = None) -> None:
        self.service, self.skill = service, skill; self.registry = registry or create_desktop_registry(service)

    def prepare(self, task: Task, payments_xlsx: str, invoice_pdfs: list[str], rules_docx: str, output_name: str = "flagged_payments.xlsx") -> tuple[ReimbursementResult, Path]:
        if task.status != TaskStatus.DRAFT: raise TaskStateError("只有 draft 任务可准备报销核对")
        result, staged = self._analyze_and_stage(task, payments_xlsx, invoice_pdfs, rules_docx, output_name)
        action = PlannedAction("deliver-flagged-payments", "spreadsheet.commit", ActionKind.WRITE, {"destination": str(self.skill.workspace.output_root / staged.name)}, "确认后交付问题支付清单", preview={"staged_path": str(staged), "analysis": asdict(result)})
        self.service.set_plan(task, "核对支付流水与发票；保守标记异常项，确认后交付问题清单。", [action])
        return result, staged

    def run(self, task: Task, payments_xlsx: str, invoice_pdfs: list[str], rules_docx: str, renderer: DocxRenderer, recipient: str, output_name: str = "flagged_payments.xlsx") -> dict[str, object]:
        """Build the complete auxiliary demo, then request one bundle approval."""
        if task.status != TaskStatus.DRAFT: raise TaskStateError("只有 draft 任务可运行完整报销核对")
        result, flagged_xlsx = self._analyze_and_stage(task, payments_xlsx, invoice_pdfs, rules_docx, output_name)
        summary, pdf, email = self.stage_summary_and_email(task, result, flagged_xlsx, renderer, recipient)
        delivery = ArtifactBundleDelivery(self.service, self.skill.workspace, self.registry)
        artifacts = [
            delivery.capture(kind="xlsx", staged_path=flagged_xlsx, final_filename=flagged_xlsx.name, summary="交付报销异常项工作簿"),
            delivery.capture(kind="pdf", staged_path=Path(pdf.staged_path), final_filename=Path(pdf.final_path).name, summary="交付经过渲染检查的报销核对摘要 PDF"),
            delivery.capture(kind="eml", staged_path=Path(email.staged_path), final_filename=Path(email.final_path).name, summary="交付明确标记为未发送的报销复核邮件草稿"),
        ]
        confirmation = self.skill.workspace.task_dir(task.task_id) / "staging" / "reimbursement_confirmation.html"
        confirmation.write_text(
            "<html><meta charset='utf-8'><body><h1>报销核对：人工确认</h1>"
            f"<p>正常 {result.matched_count} 项，待复核 {result.flagged_count} 项。</p>"
            f"<p>已应用规则：{result.applied_rules}</p>"
            "<p>确认后仅交付 XLSX、PDF 和未发送邮件草稿，不会发送邮件。</p></body></html>",
            encoding="utf-8",
        )
        self.service.set_plan(
            task,
            "读取支付流水、发票和规则文档，生成核对产物；人工确认一次后交付 XLSX、PDF 和未发送邮件草稿。",
            [
                delivery.action(action_id="reimbursement-commit-xlsx", skill="spreadsheet.commit", artifact=artifacts[0]),
                delivery.action(action_id="reimbursement-commit-pdf", skill="document.commit_pdf", artifact=artifacts[1]),
                delivery.action(action_id="reimbursement-commit-eml", skill="mail.commit_eml", artifact=artifacts[2]),
            ],
        )
        return {"task": task, "result": result, "xlsx": flagged_xlsx, "docx": summary, "pdf": Path(pdf.staged_path), "email": Path(email.staged_path), "confirmation": confirmation}

    def confirm_and_deliver(self, task: Task, *, approved: bool) -> Task:
        return ArtifactBundleDelivery(self.service, self.skill.workspace, self.registry).confirm_and_deliver(task, approved=approved)

    def _analyze_and_stage(self, task: Task, payments_xlsx: str, invoice_pdfs: list[str], rules_docx: str, output_name: str) -> tuple[ReimbursementResult, Path]:
        result = self.skill.analyze(payments_xlsx, invoice_pdfs, rules_docx)
        staged = self.skill.write_flagged_xlsx(task.task_id, output_name, result)
        self.service.trace_store.append(task.task_id, "reimbursement_analyzed", {"matched_count": result.matched_count, "flagged_count": result.flagged_count, "applied_rules": result.applied_rules, "source_hashes": result.source_hashes, "findings": [asdict(x) for x in result.findings]})
        self.service.trace_store.append(task.task_id, "reimbursement_xlsx_staged", {"path": str(staged)})
        return result, staged

    def stage_summary_and_email(self, task: Task, result: ReimbursementResult, flagged_xlsx: Path, renderer: DocxRenderer, recipient: str) -> tuple[Path, StagedPdf, StagedEmailDraft]:
        """生成可审阅 DOCX、经渲染验证的 PDF 和未发送 EML；均留在 staging。"""
        if task.status not in {TaskStatus.AWAITING_CONFIRMATION, TaskStatus.DRAFT}: raise TaskStateError("报销任务状态不允许追加交付物")
        summary = self.skill.workspace.task_dir(task.task_id) / "staging" / "reimbursement_summary.docx"
        doc = Document(); doc.add_heading("Reimbursement reconciliation summary", 0)
        doc.add_paragraph(f"Matched: {result.matched_count}; flagged: {result.flagged_count}.")
        table = doc.add_table(rows=1, cols=4); table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, ["Payment", "Invoice", "Status", "Detail"]): cell.text = value
        for item in result.findings:
            if item.status != "matched":
                row = table.add_row().cells
                for cell, value in zip(row, [item.payment_id, item.invoice_no, item.status, item.detail]): cell.text = value
        doc.save(summary)
        pdf = PdfDeliverySkill(self.skill.workspace).stage_pdf(task.task_id, summary, "reimbursement_summary.pdf", renderer)
        email = EmailDraftSkill(self.skill.workspace).stage_draft(task.task_id, "reimbursement_review.eml", [recipient], "Reimbursement reconciliation draft", f"{result.flagged_count} payments require review. This is a local draft and was not sent.", [flagged_xlsx, pdf.staged_path])
        self.service.trace_store.append(task.task_id, "reimbursement_summary_staged", {"docx": str(summary), "pdf": asdict(pdf), "email": asdict(email)})
        return summary, pdf, email
