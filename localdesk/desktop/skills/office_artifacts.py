"""受控的 DOCX 读取、PDF 交付和本地邮件草稿能力。

这些能力只处理用户授权的本地文件。所有新产物先写入任务 staging，
在交付前复核来源/附件/产物哈希；`.eml` 只是标准邮件草稿文件，模块中
没有 SMTP、Outlook 或任何发送实现。
"""

from __future__ import annotations

import hashlib
import os
import re
import shutil
from dataclasses import dataclass
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from pathlib import Path
from typing import Iterable

from docx import Document
from pypdf import PdfReader

from localdesk.desktop.docx_delivery import DocxRenderer
from localdesk.desktop.workspace import DesktopWorkspace, WorkspaceError


class OfficeArtifactError(ValueError):
    pass


@dataclass(frozen=True)
class DocxTableSummary:
    index: int
    row_count: int
    column_count: int
    rows: list[list[str]]


@dataclass(frozen=True)
class DocxInspection:
    source_path: str
    source_sha256: str
    paragraph_count: int
    headings: list[str]
    paragraphs: list[str]
    tables: list[DocxTableSummary]


@dataclass(frozen=True)
class PdfValidation:
    approved: bool
    findings: list[str]
    page_count: int
    text_page_count: int


@dataclass(frozen=True)
class StagedPdf:
    staged_path: str
    final_path: str
    source_docx: str
    source_sha256: str
    sha256: str
    validation: PdfValidation
    render_pdf_path: str
    render_image_paths: list[str]
    summary: str


@dataclass(frozen=True)
class EmailDraftValidation:
    approved: bool
    findings: list[str]
    recipient_count: int
    attachment_count: int


@dataclass(frozen=True)
class StagedEmailDraft:
    staged_path: str
    final_path: str
    sha256: str
    recipients: list[str]
    subject: str
    attachments: list[dict[str, str]]
    validation: EmailDraftValidation
    summary: str


class DocxReadSkill:
    """读取 DOCX 正文、标题层级和表格；不修改源文件。"""

    def __init__(self, workspace: DesktopWorkspace) -> None:
        self.workspace = workspace

    def inspect(self, source_path: str | Path, max_paragraphs: int = 200, max_table_rows: int = 100) -> DocxInspection:
        source = _authorized_file(self.workspace, source_path, {".docx"}, "DOCX")
        try:
            document = Document(source)
        except Exception as exc:
            raise OfficeArtifactError(f"DOCX 无法打开: {exc}") from exc
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        headings = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip() and paragraph.style and paragraph.style.name.startswith("Heading")
        ]
        tables: list[DocxTableSummary] = []
        for index, table in enumerate(document.tables, start=1):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            tables.append(DocxTableSummary(index, len(rows), len(table.columns), rows[:max_table_rows]))
        return DocxInspection(str(source), _sha256(source), len(paragraphs), headings, paragraphs[:max_paragraphs], tables)


class PdfDeliverySkill:
    """将已通过 LibreOffice 渲染的 DOCX 转为可交付 PDF。"""

    def __init__(self, workspace: DesktopWorkspace) -> None:
        self.workspace = workspace

    def stage_pdf(self, task_id: str, source_docx: str | Path, output_filename: str, renderer: DocxRenderer) -> StagedPdf:
        source = _authorized_file(self.workspace, source_docx, {".docx"}, "DOCX")
        filename = _safe_filename(output_filename, ".pdf")
        staged = self.workspace.task_dir(task_id) / "staging" / filename
        final = self.workspace.output_root / filename
        if staged.exists() or final.exists():
            raise WorkspaceError(f"禁止覆盖已有文件: {staged if staged.exists() else final}")
        source_sha256 = _sha256(source)
        render_dir = self.workspace.task_dir(task_id) / "render" / Path(filename).stem
        render = renderer.render(source, render_dir)
        if not render.approved or not render.pdf_path:
            raise OfficeArtifactError("DOCX 渲染检查未通过，不能交付正式 PDF")
        rendered_pdf = Path(render.pdf_path)
        if not rendered_pdf.is_file() or rendered_pdf.suffix.lower() != ".pdf":
            raise OfficeArtifactError("渲染器没有返回可用的 PDF 产物")
        staged.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(rendered_pdf, staged)
        validation = _validate_pdf(staged, render.page_count, render.image_paths)
        if not validation.approved:
            raise OfficeArtifactError("PDF 确定性检查未通过: " + "; ".join(validation.findings))
        return StagedPdf(
            str(staged), str(final), str(source), source_sha256, _sha256(staged), validation,
            str(rendered_pdf), list(render.image_paths), "已通过 DOCX 渲染、PDF 结构和页面检查的正式 PDF 草稿",
        )

    def validate_commit(self, draft: StagedPdf) -> None:
        staged, final, source = Path(draft.staged_path), Path(draft.final_path), Path(draft.source_docx)
        if not self.workspace.can_write_artifact(staged) or not self.workspace.is_output_artifact(final):
            raise WorkspaceError("PDF staging 或交付目标越出授权范围")
        if not staged.is_file() or _sha256(staged) != draft.sha256:
            raise WorkspaceError("PDF staging 已变化，确认失效")
        if not source.is_file() or _sha256(source) != draft.source_sha256:
            raise WorkspaceError("DOCX 来源已变化，需要重新预览和确认")
        if final.exists():
            raise WorkspaceError(f"禁止覆盖已有文件: {final}")
        if not draft.validation.approved:
            raise OfficeArtifactError("PDF 尚未通过确定性检查")

    def commit(self, draft: StagedPdf) -> None:
        self.validate_commit(draft)
        _copy_atomically(Path(draft.staged_path), Path(draft.final_path))


class EmailDraftSkill:
    """生成 RFC 兼容的本地 .eml 草稿，故意不提供发送能力。"""

    def __init__(self, workspace: DesktopWorkspace) -> None:
        self.workspace = workspace

    def stage_draft(
        self,
        task_id: str,
        output_filename: str,
        recipients: Iterable[str],
        subject: str,
        body: str,
        attachments: Iterable[str | Path] = (),
    ) -> StagedEmailDraft:
        filename = _safe_filename(output_filename, ".eml")
        staged = self.workspace.task_dir(task_id) / "staging" / filename
        final = self.workspace.output_root / filename
        if staged.exists() or final.exists():
            raise WorkspaceError(f"禁止覆盖已有文件: {staged if staged.exists() else final}")
        checked_recipients = _recipients(recipients)
        if not subject.strip() or "\r" in subject or "\n" in subject:
            raise OfficeArtifactError("邮件主题不能为空且不能包含换行")
        if "\x00" in body:
            raise OfficeArtifactError("邮件正文不能包含 NUL 字符")
        checked_attachments = [_attachment(self.workspace, item) for item in attachments]
        message = EmailMessage(policy=policy.default)
        message["To"] = ", ".join(checked_recipients)
        message["Subject"] = subject.strip()
        message["X-LocalDesk-Draft"] = "true"
        message.set_content(body)
        for item in checked_attachments:
            message.add_attachment(item["path"].read_bytes(), maintype="application", subtype="octet-stream", filename=item["name"])
        staged.parent.mkdir(parents=True, exist_ok=True)
        staged.write_bytes(message.as_bytes())
        attachment_data = [{"path": str(item["path"]), "name": item["name"], "sha256": item["sha256"]} for item in checked_attachments]
        validation = _validate_eml(staged, checked_recipients, subject.strip(), attachment_data)
        if not validation.approved:
            raise OfficeArtifactError("EML 草稿确定性检查未通过: " + "; ".join(validation.findings))
        return StagedEmailDraft(str(staged), str(final), _sha256(staged), checked_recipients, subject.strip(), attachment_data, validation, "仅本地保存、未发送的标准 .eml 邮件草稿")

    def validate_commit(self, draft: StagedEmailDraft) -> None:
        staged, final = Path(draft.staged_path), Path(draft.final_path)
        if not self.workspace.can_write_artifact(staged) or not self.workspace.is_output_artifact(final):
            raise WorkspaceError("EML staging 或交付目标越出授权范围")
        if not staged.is_file() or _sha256(staged) != draft.sha256:
            raise WorkspaceError("EML staging 已变化，确认失效")
        if final.exists():
            raise WorkspaceError(f"禁止覆盖已有文件: {final}")
        for item in draft.attachments:
            path = _authorized_file(self.workspace, item["path"], None, "附件")
            if _sha256(path) != item["sha256"]:
                raise WorkspaceError(f"邮件附件已变化，需要重新预览和确认: {path.name}")
        if not draft.validation.approved:
            raise OfficeArtifactError("EML 草稿尚未通过确定性检查")

    def commit(self, draft: StagedEmailDraft) -> None:
        self.validate_commit(draft)
        _copy_atomically(Path(draft.staged_path), Path(draft.final_path))


def _validate_pdf(path: Path, expected_pages: int, image_paths: list[str]) -> PdfValidation:
    findings: list[str] = []
    try:
        reader = PdfReader(str(path))
        pages = list(reader.pages)
        text_pages = sum(1 for page in pages if (page.extract_text() or "").strip())
    except Exception as exc:
        return PdfValidation(False, [f"PDF 无法重新打开: {exc}"], 0, 0)
    if not pages:
        findings.append("PDF 没有页面")
    if expected_pages < 1 or len(pages) != expected_pages:
        findings.append("PDF 页数与渲染检查结果不一致")
    if text_pages != len(pages):
        findings.append("PDF 存在无法提取正文的页面")
    if len(image_paths) != len(pages) or any(not Path(item).is_file() or not Path(item).stat().st_size for item in image_paths):
        findings.append("PDF 缺少可追溯的逐页渲染检查图")
    return PdfValidation(not findings, findings, len(pages), text_pages)


def _validate_eml(path: Path, recipients: list[str], subject: str, attachments: list[dict[str, str]]) -> EmailDraftValidation:
    findings: list[str] = []
    try:
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    except Exception as exc:
        return EmailDraftValidation(False, [f"EML 无法重新解析: {exc}"], 0, 0)
    parsed_to = [address.addr_spec for address in message["To"].addresses] if message["To"] else []
    if parsed_to != recipients:
        findings.append("EML 收件人与预览不一致")
    if message.get("Subject") != subject:
        findings.append("EML 主题与预览不一致")
    if message.get("Bcc"):
        findings.append("EML 不允许包含 Bcc")
    if message.get("X-LocalDesk-Draft") != "true":
        findings.append("EML 缺少本地草稿标记")
    parts = list(message.iter_attachments())
    if len(parts) != len(attachments):
        findings.append("EML 附件数量与预览不一致")
    elif [part.get_filename() for part in parts] != [item["name"] for item in attachments]:
        findings.append("EML 附件名称与预览不一致")
    return EmailDraftValidation(not findings, findings, len(parsed_to), len(parts))


def _authorized_file(workspace: DesktopWorkspace, source_path: str | Path, suffixes: set[str] | None, label: str) -> Path:
    source = workspace.resolve_path(source_path)
    if not workspace.can_read(source):
        raise WorkspaceError(f"读取路径越出授权范围: {source}")
    if not source.is_file() or source.is_symlink():
        raise OfficeArtifactError(f"{label} 必须是授权目录中的普通文件")
    if suffixes is not None and source.suffix.lower() not in suffixes:
        raise OfficeArtifactError(f"{label} 文件类型不支持: {source.suffix}")
    return source


def _attachment(workspace: DesktopWorkspace, value: str | Path) -> dict[str, object]:
    path = _authorized_file(workspace, value, None, "附件")
    return {"path": path, "name": path.name, "sha256": _sha256(path)}


def _recipients(values: Iterable[str]) -> list[str]:
    result: list[str] = []
    for raw in values:
        value = raw.strip()
        if not re.fullmatch(r"[^\s@<>]+@[^\s@<>]+\.[^\s@<>]+", value):
            raise OfficeArtifactError(f"无效收件人地址: {raw}")
        if value not in result:
            result.append(value)
    if not result:
        raise OfficeArtifactError("邮件草稿至少需要一个收件人")
    return result


def _safe_filename(filename: str, suffix: str) -> str:
    name = Path(filename).name
    if name != filename or not name or name in {".", ".."}:
        raise WorkspaceError("交付文件名必须是不含目录的文件名")
    if not name.lower().endswith(suffix):
        name += suffix
    if not re.fullmatch(r"[\w.\-\u4e00-\u9fff ]+", name):
        raise WorkspaceError("交付文件名包含不允许的字符")
    return name


def _copy_atomically(source: Path, final: Path) -> None:
    final.parent.mkdir(parents=True, exist_ok=True)
    temporary = final.with_suffix(final.suffix + ".tmp")
    shutil.copyfile(source, temporary)
    os.replace(temporary, final)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
