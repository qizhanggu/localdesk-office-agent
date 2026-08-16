"""合成或授权材料上的确定性报销核对核心；不连接财务系统、不外发。"""
from __future__ import annotations

import hashlib
import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

from localdesk.desktop.workspace import DesktopWorkspace, WorkspaceError


class ReimbursementError(ValueError): pass


@dataclass(frozen=True)
class InvoiceRecord:
    invoice_no: str
    amount: float | None
    source: str


@dataclass(frozen=True)
class PaymentFinding:
    payment_id: str
    invoice_no: str
    amount: float
    status: str
    detail: str


@dataclass(frozen=True)
class ReimbursementResult:
    findings: list[PaymentFinding]
    matched_count: int
    flagged_count: int
    source_hashes: dict[str, str]
    applied_rules: dict[str, float | None]


@dataclass(frozen=True)
class ReimbursementRules:
    max_payment_amount: float | None
    amount_tolerance: float


class ReimbursementSkill:
    """匹配 payment_id/invoice_no/amount，明确保守地标记不确定项。"""
    REQUIRED_COLUMNS = ("payment_id", "invoice_no", "amount")

    def __init__(self, workspace: DesktopWorkspace) -> None: self.workspace = workspace

    def analyze(self, payments_xlsx: str | Path, invoice_pdfs: list[str | Path], rules_docx: str | Path) -> ReimbursementResult:
        payments = self._file(payments_xlsx, ".xlsx")
        rules = self._file(rules_docx, ".docx")
        parsed_rules = self._read_rules(rules)
        invoices = [self._file(path, ".pdf") for path in invoice_pdfs]
        if not invoices: raise ReimbursementError("至少需要一张发票 PDF")
        records = self._read_invoices(invoices)
        by_no: dict[str, list[InvoiceRecord]] = {}
        for record in records: by_no.setdefault(record.invoice_no, []).append(record)
        book = load_workbook(payments, read_only=True, data_only=True)
        sheet = book.active; rows = list(sheet.iter_rows(values_only=True)); book.close()
        if not rows: raise ReimbursementError("支付流水为空")
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        missing = set(self.REQUIRED_COLUMNS) - set(headers)
        if missing: raise ReimbursementError("支付流水缺少字段: " + ", ".join(sorted(missing)))
        pos = {name: headers.index(name) for name in headers}; findings: list[PaymentFinding] = []
        for row in rows[1:]:
            payment_id, invoice_no = str(row[pos["payment_id"]] or "").strip(), str(row[pos["invoice_no"]] or "").strip()
            try: amount = float(row[pos["amount"]])
            except (TypeError, ValueError): amount = 0.0
            candidates = by_no.get(invoice_no, [])
            if not invoice_no or not candidates: status, detail = "missing_invoice", "未找到对应发票号"
            elif len(candidates) > 1: status, detail = "duplicate_invoice", f"发票号在 {len(candidates)} 个 PDF 中重复"
            elif candidates[0].amount is None: status, detail = "unconfirmed", "发票金额无法从 PDF 文本确认"
            elif abs(candidates[0].amount - amount) > parsed_rules.amount_tolerance: status, detail = "amount_mismatch", f"付款 {amount:.2f}，发票 {candidates[0].amount:.2f}"
            elif parsed_rules.max_payment_amount is not None and amount > parsed_rules.max_payment_amount: status, detail = "policy_limit_exceeded", f"付款 {amount:.2f} 超过规则上限 {parsed_rules.max_payment_amount:.2f}"
            else: status, detail = "matched", "发票号和金额一致"
            findings.append(PaymentFinding(payment_id, invoice_no, amount, status, detail))
        hashes = {str(payments): _hash(payments), str(rules): _hash(rules), **{str(path): _hash(path) for path in invoices}}
        return ReimbursementResult(
            findings,
            sum(item.status == "matched" for item in findings),
            sum(item.status != "matched" for item in findings),
            hashes,
            {
                "max_payment_amount": parsed_rules.max_payment_amount,
                "amount_tolerance": parsed_rules.amount_tolerance,
            },
        )

    def write_flagged_xlsx(self, task_id: str, filename: str, result: ReimbursementResult) -> Path:
        target = self.workspace.task_dir(task_id) / "staging" / _name(filename, ".xlsx")
        if target.exists(): raise WorkspaceError(f"禁止覆盖已有文件: {target}")
        book = Workbook(); sheet = book.active; sheet.title = "flagged_payments"
        sheet.append(["payment_id", "invoice_no", "amount", "status", "detail"])
        for item in result.findings:
            if item.status != "matched": sheet.append([item.payment_id, item.invoice_no, item.amount, item.status, item.detail])
        for cell in sheet[1]:
            font = copy(cell.font); font.bold = True; cell.font = font
        sheet.column_dimensions["A"].width = 18; sheet.column_dimensions["B"].width = 18; sheet.column_dimensions["C"].width = 14; sheet.column_dimensions["D"].width = 20; sheet.column_dimensions["E"].width = 42
        target.parent.mkdir(parents=True, exist_ok=True); book.save(target); return target

    def _read_invoices(self, paths: list[Path]) -> list[InvoiceRecord]:
        result = []
        for path in paths:
            text = "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
            no = re.search(r"invoice\s*(?:no|number|#)?\s*[:：#]?\s*([A-Za-z0-9-]+)", text, re.I)
            amount = re.search(r"(?:total|amount\s*due)\s*[:：]?\s*[$¥]?\s*([0-9]+(?:\.[0-9]{1,2})?)", text, re.I)
            if no: result.append(InvoiceRecord(no.group(1), float(amount.group(1)) if amount else None, str(path)))
        return result

    @staticmethod
    def _read_rules(path: Path) -> ReimbursementRules:
        document = Document(path)
        chunks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            chunks.extend(cell.text for row in table.rows for cell in row.cells)
        text = "\n".join(chunks)
        max_match = re.search(
            r"(?:max_payment_amount|单笔报销金额上限)\s*[:：=]\s*([0-9]+(?:\.[0-9]+)?)",
            text,
            re.I,
        )
        tolerance_match = re.search(
            r"(?:amount_tolerance|金额容差)\s*[:：=]\s*([0-9]+(?:\.[0-9]+)?)",
            text,
            re.I,
        )
        if not max_match and not tolerance_match:
            raise ReimbursementError(
                "规则文档未包含可识别规则；请使用 max_payment_amount / 单笔报销金额上限，"
                "或 amount_tolerance / 金额容差。"
            )
        maximum = float(max_match.group(1)) if max_match else None
        tolerance = float(tolerance_match.group(1)) if tolerance_match else 0.005
        if maximum is not None and maximum <= 0:
            raise ReimbursementError("单笔报销金额上限必须大于 0。")
        if tolerance < 0:
            raise ReimbursementError("金额容差不能小于 0。")
        return ReimbursementRules(maximum, tolerance)

    def _file(self, value: str | Path, suffix: str) -> Path:
        path = self.workspace.resolve_path(value)
        if not self.workspace.can_read(path): raise WorkspaceError(f"读取路径越出授权范围: {path}")
        if not path.is_file() or path.is_symlink() or path.suffix.lower() != suffix: raise ReimbursementError(f"需要授权的普通 {suffix} 文件: {path}")
        return path

def _name(value: str, suffix: str) -> str:
    name = Path(value).name
    if name != value or not re.fullmatch(r"[\w.\-\u4e00-\u9fff ]+", name): raise WorkspaceError("输出文件名不安全")
    return name if name.lower().endswith(suffix) else name + suffix
def _hash(path: Path) -> str: return hashlib.sha256(path.read_bytes()).hexdigest()
