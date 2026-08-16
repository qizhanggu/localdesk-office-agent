"""受控的 CSV/XLSX 结构化读写能力。

该模块只处理工作簿数据，不驱动 Excel GUI。所有副作用先写入 task
staging，交付时重新检查输入哈希、staging 哈希与目标冲突。
"""

from __future__ import annotations

import csv
import hashlib
import os
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from localdesk.desktop.workspace import DesktopWorkspace, WorkspaceError


SUPPORTED_SUFFIXES = frozenset({".csv", ".xlsx"})


class SpreadsheetError(ValueError):
    pass


@dataclass(frozen=True)
class SortRule:
    column: str
    ascending: bool = True


@dataclass(frozen=True)
class FilterRule:
    column: str
    equals: Any


@dataclass(frozen=True)
class ColumnUpdate:
    column: str
    value: Any
    when_column: str | None = None
    when_equals: Any = None


@dataclass(frozen=True)
class SpreadsheetTransform:
    sheet_name: str | None = None
    filters: tuple[FilterRule, ...] = ()
    sort_rules: tuple[SortRule, ...] = ()
    dedupe_columns: tuple[str, ...] = ()
    column_updates: tuple[ColumnUpdate, ...] = ()


@dataclass(frozen=True)
class SheetSummary:
    name: str
    headers: list[str]
    row_count: int
    column_count: int
    sample_rows: list[list[Any]]


@dataclass(frozen=True)
class SpreadsheetInspection:
    source_path: str
    source_sha256: str
    file_type: str
    sheets: list[SheetSummary]


@dataclass(frozen=True)
class SpreadsheetValidation:
    approved: bool
    findings: list[str]
    expected_row_count: int
    actual_row_count: int


@dataclass(frozen=True)
class StagedSpreadsheet:
    staged_path: str
    final_path: str
    source_path: str
    source_sha256: str
    sha256: str
    summary: dict[str, Any]
    validation: SpreadsheetValidation


@dataclass
class _Table:
    headers: list[str]
    rows: list[list[Any]]


class SpreadsheetSkill:
    def __init__(self, workspace: DesktopWorkspace) -> None:
        self.workspace = workspace

    def inspect(self, source_path: str | Path, sample_rows: int = 20) -> SpreadsheetInspection:
        source = self._authorized_source(source_path)
        if source.suffix.lower() == ".csv":
            table = self._read_csv(source)
            sheets = [self._summary("Sheet1", table, sample_rows)]
        else:
            workbook = load_workbook(source, read_only=True, data_only=False)
            sheets = [
                self._summary(sheet.title, self._read_sheet(sheet), sample_rows)
                for sheet in workbook.worksheets
            ]
            workbook.close()
        return SpreadsheetInspection(str(source), _sha256(source), source.suffix.lower(), sheets)

    def read_cells(self, source_path: str | Path, sheet_name: str | None = None, max_rows: int = 100) -> SheetSummary:
        inspection = self.inspect(source_path, sample_rows=max_rows)
        if sheet_name is None:
            return inspection.sheets[0]
        for sheet in inspection.sheets:
            if sheet.name == sheet_name:
                return sheet
        raise SpreadsheetError(f"工作表不存在: {sheet_name}")

    def stage_transform(self, task_id: str, source_path: str | Path, output_filename: str, transform: SpreadsheetTransform) -> StagedSpreadsheet:
        source = self._authorized_source(source_path)
        filename = _safe_filename(output_filename)
        staged = self.workspace.task_dir(task_id) / "staging" / filename
        final = self.workspace.output_root / filename
        if staged.exists() or final.exists():
            raise WorkspaceError(f"禁止覆盖已有文件: {staged if staged.exists() else final}")
        source_hash = _sha256(source)
        if source.suffix.lower() == ".csv":
            if filename.lower().endswith(".xlsx"):
                raise SpreadsheetError("CSV 输入当前只能交付为新的 CSV 文件")
            table = self._read_csv(source)
            transformed, summary = _transform_table(table, transform)
            staged.parent.mkdir(parents=True, exist_ok=True)
            self._write_csv(staged, transformed)
            actual = self._read_csv(staged)
        else:
            if filename.lower().endswith(".csv"):
                raise SpreadsheetError("XLSX 输入当前只能交付为新的 XLSX 文件")
            summary, actual = self._write_xlsx_transform(source, staged, transform)
            transformed = actual
        validation = self._validate(source, source_hash, transformed, actual, summary)
        if not validation.approved:
            raise SpreadsheetError("工作簿确定性校验失败: " + "; ".join(validation.findings))
        return StagedSpreadsheet(str(staged), str(final), str(source), source_hash, _sha256(staged), summary, validation)

    def validate_commit(self, draft: StagedSpreadsheet) -> None:
        staged, final, source = Path(draft.staged_path), Path(draft.final_path), Path(draft.source_path)
        if not self.workspace.can_write_artifact(staged) or not self.workspace.can_write_artifact(final):
            raise WorkspaceError("工作簿 staging 或交付目标越出授权范围")
        if not staged.is_file() or _sha256(staged) != draft.sha256:
            raise WorkspaceError("工作簿 staging 已变化，确认失效")
        if not source.is_file() or _sha256(source) != draft.source_sha256:
            raise WorkspaceError("输入工作簿已变化，需要重新预览和确认")
        if final.exists():
            raise WorkspaceError(f"禁止覆盖已有文件: {final}")
        if not draft.validation.approved:
            raise SpreadsheetError("未通过确定性工作簿校验")

    def commit(self, draft: StagedSpreadsheet) -> None:
        self.validate_commit(draft)
        final = Path(draft.final_path)
        final.parent.mkdir(parents=True, exist_ok=True)
        temporary = final.with_suffix(final.suffix + ".tmp")
        temporary.write_bytes(Path(draft.staged_path).read_bytes())
        os.replace(temporary, final)

    def _authorized_source(self, source_path: str | Path) -> Path:
        source = self.workspace.resolve_path(source_path)
        if not self.workspace.can_read(source):
            raise WorkspaceError(f"读取路径越出授权范围: {source}")
        if source.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise SpreadsheetError("仅支持 .csv 和 .xlsx 文件")
        if not source.is_file() or source.is_symlink():
            raise SpreadsheetError("输入必须是授权目录中的普通工作簿文件")
        return source

    @staticmethod
    def _read_csv(path: Path) -> _Table:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        if not rows:
            raise SpreadsheetError("CSV 为空，缺少表头")
        return _Table(_headers(rows[0]), [list(row) + [""] * (len(rows[0]) - len(row)) for row in rows[1:]])

    @staticmethod
    def _write_csv(path: Path, table: _Table) -> None:
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(table.headers)
            writer.writerows(table.rows)

    @staticmethod
    def _read_sheet(sheet: Any) -> _Table:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise SpreadsheetError(f"工作表为空，缺少表头: {sheet.title}")
        headers = _headers(list(rows[0]))
        width = len(headers)
        return _Table(headers, [list(row[:width]) + [None] * (width - len(row[:width])) for row in rows[1:]])

    @staticmethod
    def _summary(name: str, table: _Table, sample_rows: int) -> SheetSummary:
        return SheetSummary(name, table.headers, len(table.rows), len(table.headers), table.rows[:sample_rows])

    def _write_xlsx_transform(self, source: Path, staged: Path, transform: SpreadsheetTransform) -> tuple[dict[str, Any], _Table]:
        workbook = load_workbook(source, data_only=False)
        target_name = transform.sheet_name or workbook.sheetnames[0]
        if target_name not in workbook.sheetnames:
            workbook.close()
            raise SpreadsheetError(f"工作表不存在: {target_name}")
        sheet = workbook[target_name]
        original = self._read_sheet(sheet)
        transformed, summary = _transform_table(original, transform)
        old_row_count = max(sheet.max_row - 1, 0)
        if old_row_count:
            sheet.delete_rows(2, old_row_count)
        for column, header in enumerate(transformed.headers, start=1):
            sheet.cell(1, column).value = header
        for row_index, row in enumerate(transformed.rows, start=2):
            for column, value in enumerate(row, start=1):
                sheet.cell(row_index, column).value = value
        staged.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(staged)
        workbook.close()
        verify = load_workbook(staged, read_only=True, data_only=False)
        actual = self._read_sheet(verify[target_name])
        verify.close()
        summary["sheet_name"] = target_name
        return summary, actual

    @staticmethod
    def _validate(source: Path, expected_source_hash: str, expected: _Table, actual: _Table, summary: dict[str, Any]) -> SpreadsheetValidation:
        findings: list[str] = []
        if _sha256(source) != expected_source_hash:
            findings.append("输入文件哈希在处理期间变化")
        if expected.headers != actual.headers:
            findings.append("输出表头与计划不一致")
        if expected.rows != actual.rows:
            findings.append("输出单元格值或数据类型与确定性计划不一致")
        return SpreadsheetValidation(not findings, findings, int(summary["output_row_count"]), len(actual.rows))


def _headers(values: list[Any]) -> list[str]:
    headers = [str(value).strip() if value is not None else "" for value in values]
    if not headers or any(not header for header in headers):
        raise SpreadsheetError("表头不能为空")
    if len(set(headers)) != len(headers):
        raise SpreadsheetError("表头不能重复")
    return headers


def _transform_table(table: _Table, transform: SpreadsheetTransform) -> tuple[_Table, dict[str, Any]]:
    headers, rows = list(table.headers), [list(row) for row in table.rows]
    index = {header: position for position, header in enumerate(headers)}
    _require_columns(index, [rule.column for rule in transform.filters])
    for rule in transform.filters:
        rows = [row for row in rows if row[index[rule.column]] == rule.equals]
    before_dedupe = len(rows)
    if transform.dedupe_columns:
        _require_columns(index, transform.dedupe_columns)
        seen: set[tuple[Any, ...]] = set(); unique: list[list[Any]] = []
        for row in rows:
            key = tuple(row[index[column]] for column in transform.dedupe_columns)
            if key not in seen:
                seen.add(key); unique.append(row)
        rows = unique
    _require_columns(index, [rule.column for rule in transform.sort_rules])
    for rule in reversed(transform.sort_rules):
        rows.sort(key=lambda row, position=index[rule.column]: _sort_value(row[position]), reverse=not rule.ascending)
    updates: list[str] = []
    for update in transform.column_updates:
        if update.column not in index:
            index[update.column] = len(headers); headers.append(update.column)
            for row in rows: row.append(None)
        position = index[update.column]
        if update.when_column is not None:
            _require_columns(index, [update.when_column])
            match_position = index[update.when_column]
            selected = [row for row in rows if row[match_position] == update.when_equals]
        else:
            selected = rows
        for row in selected: row[position] = update.value
        updates.append(f"{update.column}:{len(selected)}")
    summary = {
        "input_row_count": len(table.rows), "output_row_count": len(rows),
        "filtered_row_count": before_dedupe, "deduped_removed": before_dedupe - len(rows),
        "filters": [asdict(item) for item in transform.filters],
        "sort_rules": [asdict(item) for item in transform.sort_rules],
        "dedupe_columns": list(transform.dedupe_columns), "column_updates": updates,
    }
    return _Table(headers, rows), summary


def _require_columns(index: dict[str, int], columns: Any) -> None:
    missing = [column for column in columns if column not in index]
    if missing:
        raise SpreadsheetError("找不到列: " + ", ".join(missing))


def _sort_value(value: Any) -> tuple[int, Any]:
    if value is None or value == "": return (0, "")
    if isinstance(value, bool): return (1, int(value))
    if isinstance(value, (int, float)): return (2, value)
    return (3, str(value).casefold())


def _safe_filename(filename: str) -> str:
    name = Path(filename).name
    if name != filename or not name or name in {".", ".."}:
        raise WorkspaceError("工作簿文件名必须是不含目录的文件名")
    if Path(name).suffix.lower() not in SUPPORTED_SUFFIXES:
        raise SpreadsheetError("输出仅支持 .csv 或 .xlsx")
    if not re.fullmatch(r"[\w.\-\u4e00-\u9fff ]+", name):
        raise WorkspaceError("工作簿文件名包含不允许的字符")
    return name


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
