from __future__ import annotations

from email import policy
from email.parser import BytesParser
from pathlib import Path

import fitz
import pytest
from docx import Document

from localdesk.desktop.docx_delivery import RenderCheck
from localdesk.desktop.office_artifact_workflow import OfficeArtifactWorkflow
from localdesk.desktop.policy import DesktopPolicyGuard
from localdesk.desktop.service import DesktopTaskService
from localdesk.desktop.skills.knowledge import KnowledgeSkill
from localdesk.desktop.skills.office_artifacts import DocxReadSkill, EmailDraftSkill, OfficeArtifactError
from localdesk.desktop.trace_store import TaskTraceStore
from localdesk.desktop.workspace import DesktopWorkspace, WorkspaceConfig, WorkspaceError


@pytest.fixture
def workspace(tmp_path: Path) -> DesktopWorkspace:
    sources, output, tasks = tmp_path / "sources", tmp_path / "output", tmp_path / "tasks"
    for directory in (sources, output, tasks):
        directory.mkdir()
    return DesktopWorkspace(WorkspaceConfig(read_roots=[sources], output_root=output, task_root=tasks))


@pytest.fixture
def workflow(workspace: DesktopWorkspace) -> OfficeArtifactWorkflow:
    service = DesktopTaskService(DesktopPolicyGuard(workspace), TaskTraceStore(workspace))
    return OfficeArtifactWorkflow(service, workspace)


class TestPdfRenderer:
    def render(self, _docx_path: Path, output_dir: Path) -> RenderCheck:
        output_dir.mkdir(parents=True, exist_ok=True)
        pdf_path = output_dir / "rendered.pdf"
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "LocalDesk formal delivery")
        pdf.save(pdf_path)
        pdf.close()
        image_path = output_dir / "page-1.png"
        image_path.write_bytes(b"PNG fixture")
        return RenderCheck(True, [], str(pdf_path), [str(image_path)], 1)


def _source_docx(workspace: DesktopWorkspace) -> Path:
    path = workspace.read_roots[0] / "source.docx"
    document = Document()
    document.add_heading("Expense summary", level=1)
    document.add_paragraph("Payment records were matched.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Amount", "Status"
    table.cell(1, 0).text, table.cell(1, 1).text = "100", "Matched"
    document.save(path)
    return path


def test_docx_read_supports_paragraphs_headings_and_tables(workspace: DesktopWorkspace) -> None:
    source = _source_docx(workspace)
    inspection = DocxReadSkill(workspace).inspect(source)

    assert inspection.paragraph_count == 2
    assert inspection.headings == ["Expense summary"]
    assert inspection.tables[0].rows[1] == ["100", "Matched"]

    knowledge = KnowledgeSkill(workspace)
    assert knowledge.index() == 4
    hit = knowledge.search("100")
    assert hit[0].locator == "table 1 row 2"


def test_pdf_is_staged_checked_then_confirmed(workspace: DesktopWorkspace, workflow: OfficeArtifactWorkflow) -> None:
    source = _source_docx(workspace)
    task = workflow.service.create_task("deliver checked PDF")
    draft = workflow.prepare_pdf(task, str(source), "expense-summary.pdf", TestPdfRenderer())

    assert task.status.value == "awaiting_confirmation"
    assert Path(draft.staged_path).is_file()
    assert not Path(draft.final_path).exists()
    assert draft.validation.approved and draft.validation.page_count == 1
    assert "pdf_staged" in [event["event_type"] for event in workflow.service.trace_store.load_events(task.task_id)]

    workflow.confirm_and_deliver(task, approved=True)
    assert task.status.value == "succeeded"
    assert Path(draft.final_path).is_file()
    assert task.artifacts[0].kind == "pdf"


def test_pdf_refuses_changed_docx_before_delivery(workspace: DesktopWorkspace, workflow: OfficeArtifactWorkflow) -> None:
    source = _source_docx(workspace)
    task = workflow.service.create_task("deliver checked PDF")
    draft = workflow.prepare_pdf(task, str(source), "expense-summary.pdf", TestPdfRenderer())
    source.write_bytes(b"changed")

    with pytest.raises(WorkspaceError, match="来源已变化"):
        workflow.confirm_and_deliver(task, approved=True)
    assert task.status.value == "failed"
    assert not Path(draft.final_path).exists()


def test_eml_draft_is_auto_delivered_but_never_sent(workspace: DesktopWorkspace, workflow: OfficeArtifactWorkflow) -> None:
    attachment = workspace.read_roots[0] / "summary.pdf"
    attachment.write_bytes(b"synthetic attachment")
    task = workflow.service.create_task("draft reimbursement email")
    draft = workflow.prepare_email(
        task, "reimbursement.eml", ["finance@example.com"], "Reimbursement summary", "Please review the attached summary.", [str(attachment)],
    )

    assert task.status.value == "succeeded"
    delivered = Path(draft.final_path)
    message = BytesParser(policy=policy.default).parsebytes(delivered.read_bytes())
    assert message["To"] == "finance@example.com"
    assert message["X-LocalDesk-Draft"] == "true"
    assert [part.get_filename() for part in message.iter_attachments()] == ["summary.pdf"]
    event_types = [event["event_type"] for event in workflow.service.trace_store.load_events(task.task_id)]
    assert "low_risk_auto_approved" in event_types
    assert all("send" not in event_type for event_type in event_types)


def test_eml_rejects_header_injection_and_attachment_tampering(workspace: DesktopWorkspace, workflow: OfficeArtifactWorkflow) -> None:
    attachment = workspace.read_roots[0] / "summary.pdf"
    attachment.write_bytes(b"one")
    task = workflow.service.create_task("draft mail")
    with pytest.raises(OfficeArtifactError, match="主题不能为空"):
        workflow.prepare_email(task, "draft.eml", ["finance@example.com"], "ok\r\nBcc: bad@example.com", "body")

    task = workflow.service.create_task("draft mail")
    # A symlink is rejected even when it points into an authorized source root.
    link = workspace.read_roots[0] / "link.pdf"
    try:
        link.symlink_to(attachment)
    except OSError:
        pytest.skip("symlinks unavailable in this Windows test environment")
    with pytest.raises(OfficeArtifactError, match="普通文件"):
        workflow.prepare_email(task, "draft.eml", ["finance@example.com"], "subject", "body", [str(link)])


def test_eml_commit_refuses_changed_attachment(workspace: DesktopWorkspace, workflow: OfficeArtifactWorkflow) -> None:
    attachment = workspace.read_roots[0] / "summary.pdf"
    attachment.write_bytes(b"original")
    task = workflow.service.create_task("stage local mail")
    draft = EmailDraftSkill(workspace).stage_draft(
        task.task_id, "draft.eml", ["finance@example.com"], "subject", "body", [str(attachment)],
    )
    attachment.write_bytes(b"changed")

    with pytest.raises(WorkspaceError, match="附件已变化"):
        EmailDraftSkill(workspace).validate_commit(draft)
