from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from localdesk.desktop.policy import DesktopPolicyGuard
from localdesk.desktop.service import DesktopTaskService
from localdesk.desktop.spreadsheet_workflow import SpreadsheetWorkflow
from localdesk.desktop.skills.spreadsheet import (
    ColumnUpdate,
    FilterRule,
    SortRule,
    SpreadsheetError,
    SpreadsheetSkill,
    SpreadsheetTransform,
)
from localdesk.desktop.trace_store import TaskTraceStore
from localdesk.desktop.workspace import DesktopWorkspace, WorkspaceConfig, WorkspaceError


@pytest.fixture
def workspace(tmp_path: Path) -> DesktopWorkspace:
    sources, output, tasks = tmp_path / "sources", tmp_path / "output", tmp_path / "tasks"
    for directory in (sources, output, tasks): directory.mkdir()
    return DesktopWorkspace(WorkspaceConfig(read_roots=[sources], output_root=output, task_root=tasks))


def make_book(path: Path) -> None:
    workbook = Workbook(); payments = workbook.active; payments.title = "支付流水"
    payments.append(["供应商", "金额", "备注"])
    payments.append(["乙公司", 10, "first"])
    payments.append(["甲公司", 10, "second"])
    payments.append(["乙公司", 10, "duplicate"])
    payments.append(["甲公司", 2, None])
    rules = workbook.create_sheet("规则")
    rules.append(["规则", "阈值"]); rules.append(["无需覆盖输入", 500])
    workbook.save(path)


def test_inspect_reads_chinese_headers_empty_cells_and_multiple_sheets(workspace: DesktopWorkspace) -> None:
    source = workspace.read_roots[0] / "支付流水.xlsx"; make_book(source)
    inspection = SpreadsheetSkill(workspace).inspect(source)

    assert [sheet.name for sheet in inspection.sheets] == ["支付流水", "规则"]
    payments = inspection.sheets[0]
    assert payments.headers == ["供应商", "金额", "备注"]
    assert payments.row_count == 4 and payments.column_count == 3
    assert payments.sample_rows[-1] == ["甲公司", 2, None]


def test_xlsx_transform_preserves_number_type_stable_sort_and_dedupes(workspace: DesktopWorkspace) -> None:
    source = workspace.read_roots[0] / "payments.xlsx"; make_book(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    skill = SpreadsheetSkill(workspace)
    transform = SpreadsheetTransform(
        sheet_name="支付流水",
        sort_rules=(SortRule("金额"),),
        dedupe_columns=("供应商", "金额"),
        column_updates=(ColumnUpdate("状态", "待核对", "金额", 10),),
    )
    draft = skill.stage_transform("task-1", source, "flagged_payments.xlsx", transform)
    book = load_workbook(draft.staged_path, data_only=False)
    sheet = book["支付流水"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))

    assert sheet.cell(1, 4).value == "状态"
    assert rows == [("甲公司", 2, None, None), ("乙公司", 10, "first", "待核对"), ("甲公司", 10, "second", "待核对")]
    assert isinstance(rows[0][1], int)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert draft.validation.approved and draft.summary["deduped_removed"] == 1


def test_transform_keeps_unmodified_worksheets(workspace: DesktopWorkspace) -> None:
    source = workspace.read_roots[0] / "multi.xlsx"; make_book(source)
    draft = SpreadsheetSkill(workspace).stage_transform(
        "task-2", source, "result.xlsx", SpreadsheetTransform(sheet_name="支付流水", filters=(FilterRule("供应商", "甲公司"),))
    )
    workbook = load_workbook(draft.staged_path, data_only=False)
    assert workbook["规则"]["A2"].value == "无需覆盖输入"
    assert workbook["规则"]["B2"].value == 500


def test_csv_filter_and_update_are_written_as_new_file(workspace: DesktopWorkspace) -> None:
    source = workspace.read_roots[0] / "名单.csv"
    source.write_text("姓名,部门\n张三,研发\n李四,产品\n张三,研发\n", encoding="utf-8")
    draft = SpreadsheetSkill(workspace).stage_transform(
        "task-3", source, "研发名单.csv",
        SpreadsheetTransform(filters=(FilterRule("部门", "研发"),), dedupe_columns=("姓名",), column_updates=(ColumnUpdate("状态", "已筛选"),)),
    )
    output = Path(draft.staged_path).read_text(encoding="utf-8")
    assert "姓名,部门,状态" in output and output.count("张三") == 1


def test_source_is_not_overwritten_and_existing_delivery_is_rejected(workspace: DesktopWorkspace) -> None:
    source = workspace.read_roots[0] / "source.xlsx"; make_book(source)
    source_bytes = source.read_bytes()
    (workspace.output_root / "result.xlsx").write_bytes(b"keep")
    with pytest.raises(WorkspaceError, match="禁止覆盖"):
        SpreadsheetSkill(workspace).stage_transform("task-4", source, "result.xlsx", SpreadsheetTransform())
    assert source.read_bytes() == source_bytes


def test_rejects_path_escape_and_wrong_extension(workspace: DesktopWorkspace, tmp_path: Path) -> None:
    source = workspace.read_roots[0] / "source.xlsx"; make_book(source)
    skill = SpreadsheetSkill(workspace)
    with pytest.raises(WorkspaceError, match="越出授权范围"):
        skill.inspect(tmp_path / "outside.xlsx")
    with pytest.raises(SpreadsheetError, match="仅支持"):
        skill.stage_transform("task-5", source, "result.txt", SpreadsheetTransform())
    with pytest.raises(WorkspaceError, match="不含目录"):
        skill.stage_transform("task-5", source, "../result.xlsx", SpreadsheetTransform())


def test_workflow_requires_confirmation_and_records_trace(workspace: DesktopWorkspace) -> None:
    source = workspace.read_roots[0] / "source.xlsx"; make_book(source)
    service = DesktopTaskService(DesktopPolicyGuard(workspace), TaskTraceStore(workspace))
    workflow = SpreadsheetWorkflow(service, SpreadsheetSkill(workspace))
    task = service.create_task("整理支付流水")
    draft = workflow.prepare(task, str(source), "result.xlsx", SpreadsheetTransform(sheet_name="支付流水", sort_rules=(SortRule("金额"),)))

    assert task.status.value == "awaiting_confirmation"
    assert Path(draft.staged_path).exists() and not Path(draft.final_path).exists()
    workflow.confirm_and_deliver(task, approved=True)
    assert task.status.value == "succeeded" and Path(draft.final_path).exists()
    event_types = {event["event_type"] for event in service.trace_store.load_events(task.task_id)}
    assert {"spreadsheet_staged", "artifact_committed", "tool_verified"}.issubset(event_types)


def test_input_hash_change_blocks_commit(workspace: DesktopWorkspace) -> None:
    source = workspace.read_roots[0] / "source.xlsx"; make_book(source)
    skill = SpreadsheetSkill(workspace)
    draft = skill.stage_transform("task-6", source, "result.xlsx", SpreadsheetTransform())
    source.write_bytes(source.read_bytes() + b"changed")
    with pytest.raises(WorkspaceError, match="输入工作簿已变化"):
        skill.validate_commit(draft)
